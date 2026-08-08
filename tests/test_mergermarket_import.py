import io
import unittest

import openpyxl

from mergermarket_import import parse_mergermarket_export
from pipeline_core_v2 import load_template_workbook, write_transaction_comps_workbook


SAMPLE_CSV = b"""Mergermarket Deal ID,Announcement Date,Target,Acquirer,Seller,Deal Value USD mn,Stake Acquired,Enterprise Value USD mn,Target EBITDA USD mn,Deal Description,Target Sector,Target Geography,Deal URL
MM-1,2025-01-15,Alpha Safety,Buyer Group,Founder,120,75%,150,15,Industrial safety equipment manufacturer,Industrials,India,https://example.com/mm-1
MM-1,2025-01-15,Alpha Safety,Buyer Group,Founder,120,75%,150,15,Duplicate row,Industrials,India,https://example.com/mm-1
MM-2,2024-06-30,Beta Software,Tech Buyer,,50,100%,60,0,Unrelated software provider,Technology,United States,https://example.com/mm-2
"""


class MergermarketImportTests(unittest.TestCase):
    def test_csv_import_maps_fields_and_deduplicates_by_deal_id(self):
        imported = parse_mergermarket_export("deals.csv", SAMPLE_CSV)
        self.assertEqual(len(imported.deals), 2)
        self.assertEqual(imported.deals[0]["deal_id"], "MM-1")
        self.assertEqual(imported.deals[0]["deal_percent"], 0.75)
        self.assertEqual(imported.deals[0]["enterprise_value_usd_mn"], 150.0)
        self.assertTrue(any("duplicate" in warning.lower() for warning in imported.warnings))

    def test_transaction_workbook_preserves_raw_data_and_formula_driven_multiple(self):
        imported = parse_mergermarket_export("deals.csv", SAMPLE_CSV)
        imported.deals[0]["relevance_tier"] = "Core"
        imported.deals[0]["relevance_rationale"] = "Direct product overlap"
        imported.deals[1]["relevance_tier"] = "Excluded"
        imported.deals[1]["relevance_rationale"] = "Unrelated software"

        workbook = load_template_workbook()
        write_transaction_comps_workbook(workbook, imported, "Subject Co", "Domestic and global")
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        checked = openpyxl.load_workbook(output, data_only=False)

        self.assertIn("MM Raw Export", checked.sheetnames)
        self.assertIn("Final Comps", checked.sheetnames)
        self.assertIn("Excluded", checked.sheetnames)
        self.assertIn("QA", checked.sheetnames)
        self.assertEqual(checked["Final Comps"]["B2"].value, "Alpha Safety")
        self.assertEqual(checked["Final Comps"]["G2"].value, '=IFERROR(O2/Q2,"")')
        self.assertEqual(checked["Final Comps"]["S2"].value, "Calculated")


if __name__ == "__main__":
    unittest.main()
