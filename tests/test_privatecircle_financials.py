import os
import unittest
from unittest.mock import patch

import openpyxl

from pipeline_core_v2 import (
    _privatecircle_preferred_source,
    privatecircle_amount_divisor,
    pull_financial_statements,
    write_financials_sheet,
)


def _statement(items, notes=None):
    return {"items": items, "notes": notes or ["All values are in INR mn"]}


class PrivateCircleFinancialTests(unittest.TestCase):
    def test_documented_inr_mn_values_are_not_divided_again(self):
        payload = _statement([{"fiscal_year": 2025, "total_revenue": 1250.5}])
        self.assertEqual(privatecircle_amount_divisor(payload), 1.0)

    def test_legacy_rupee_scale_values_without_notes_are_detected(self):
        payload = {"items": [{"fiscal_year": 2025, "total_revenue": 1_250_000_000}]}
        self.assertEqual(privatecircle_amount_divisor(payload), 1_000_000.0)

    def test_explicit_unit_override_handles_legacy_api_accounts(self):
        payload = _statement([{"fiscal_year": 2025, "total_revenue": 1_250_000_000}])
        with patch.dict(os.environ, {"PC_FINANCIAL_INPUT_UNIT": "inr"}):
            self.assertEqual(privatecircle_amount_divisor(payload), 1_000_000.0)

    def test_listing_status_selects_the_documented_source(self):
        self.assertEqual(_privatecircle_preferred_source({"listing_status": "Listed"}), "exchange_filings")
        self.assertEqual(_privatecircle_preferred_source({"listing_status": "Unlisted"}), "mca")

    def test_consolidated_falls_back_to_standalone_and_pulls_all_statements(self):
        calls = []

        def fake_get(_encrypted_id, suffix, _token, params=None):
            calls.append((suffix, dict(params or {})))
            if suffix == "overview":
                return {"items": [{"name": "Example Limited", "listing_status": "Unlisted"}]}
            if suffix == "income-statement" and params.get("xbrl_type_tag") == "consolidated":
                return _statement([])
            return _statement([{"fiscal_year": 2025, "xbrl_type_tag": "Standalone"}])

        with patch("pipeline_core_v2.pc_get_company_endpoint", side_effect=fake_get):
            result = pull_financial_statements("encrypted", "token")

        income_calls = [params for suffix, params in calls if suffix == "income-statement"]
        self.assertEqual(income_calls[0]["source"], "mca")
        self.assertEqual(income_calls[0]["xbrl_type_tag"], "consolidated")
        self.assertEqual(income_calls[1]["xbrl_type_tag"], "standalone")
        self.assertEqual(result["_meta"]["income_statement"]["xbrl_type_tag"], "standalone")
        self.assertTrue(result["cash_flow"]["items"])

    def test_workbook_preserves_inr_mn_amounts_and_includes_cash_flow(self):
        financial_package = {
            "overview": {
                "items": [
                    {
                        "name": "Example Limited",
                        "true_cid": "U12345MH2020PLC123456",
                        "listing_status": "Unlisted",
                    }
                ]
            },
            "income_statement": _statement(
                [
                    {
                        "fiscal_year": 2025,
                        "revenue_from_operations": 1250.5,
                        "other_income": 10.0,
                        "operating_ebitda": 200.0,
                        "pat": 100.0,
                    }
                ]
            ),
            "balance_sheet": _statement(
                [{"fiscal_year": 2025, "bs_total_assets": 900.0, "bs_total_liabilities": 900.0}]
            ),
            "cash_flow": _statement(
                [{"fiscal_year": 2025, "net_cashflows_from_operatng_activts": 140.0}]
            ),
            "basic_financial_ratios": {
                "items": [{"fiscal_year": 2025, "roce_percent": 18.0}],
                "notes": [],
            },
            "_meta": {
                name: {"source": "mca", "xbrl_type_tag": "standalone", "item_count": 1}
                for name in (
                    "income_statement",
                    "balance_sheet",
                    "cash_flow",
                    "basic_financial_ratios",
                )
            },
        }
        workbook = openpyxl.Workbook()
        workbook.active.title = "Financials"

        with patch("pipeline_core_v2.pull_financial_statements", return_value=financial_package):
            write_financials_sheet(workbook, "encrypted", "token")

        sheet = workbook["Financials"]
        labels = {sheet.cell(row, 1).value: row for row in range(1, sheet.max_row + 1)}
        self.assertEqual(sheet.cell(labels["Revenue from Operations"], 2).value, 1250.5)
        self.assertEqual(sheet.cell(labels["Total Revenue"], 2).value, 1260.5)
        self.assertEqual(sheet.cell(labels["Cash Flow from Operations"], 2).value, 140.0)
        self.assertIn("already INR mn", sheet["B6"].value)


if __name__ == "__main__":
    unittest.main()
