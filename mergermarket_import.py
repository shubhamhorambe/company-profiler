"""Parse authorised Mergermarket exports without depending on a fixed layout."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Sequence, Tuple

import openpyxl


def _normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


HEADER_ALIASES = {
    "deal_id": {"deal id", "mergermarket deal id", "deal number", "deal no", "id"},
    "announced_date": {"announced date", "announcement date", "date announced", "announcement"},
    "target": {"target", "target company", "target name", "asset target"},
    "acquirer": {"acquirer", "buyer", "bidder", "acquiring company", "acquiror"},
    "seller": {"seller", "vendor", "divestor", "selling shareholder"},
    "deal_value_original": {
        "deal value",
        "transaction value",
        "consideration",
        "deal value original",
        "value",
    },
    "deal_value_usd_mn": {
        "deal value usd mn",
        "deal value usd m",
        "deal value usdm",
        "deal value usd million",
        "transaction value usd mn",
    },
    "currency": {"currency", "deal currency", "value currency"},
    "deal_percent": {"deal percent", "deal percentage", "stake acquired", "percent acquired", "stake"},
    "deal_status": {"deal status", "status", "transaction status"},
    "deal_type": {"deal type", "transaction type", "type"},
    "description": {"description", "deal description", "transaction description", "synopsis", "deal synopsis"},
    "target_sector": {"target sector", "sector", "subsector", "target subsector"},
    "target_geography": {"target geography", "geography", "target country", "country"},
    "enterprise_value_usd_mn": {
        "enterprise value usd mn",
        "enterprise value usd m",
        "enterprise value usdm",
        "enterprise value usd million",
        "ev usd mn",
    },
    "revenue_usd_mn": {"revenue usd mn", "sales usd mn", "target revenue usd mn"},
    "ebitda_usd_mn": {"ebitda usd mn", "target ebitda usd mn", "ltm ebitda usd mn"},
    "ev_ebitda_reported": {"ev ebitda", "ev ebitda x", "enterprise value ebitda", "ev ebitda multiple"},
    "ev_revenue_reported": {"ev revenue", "ev sales", "ev revenue x", "ev sales multiple"},
    "profile_url": {"deal url", "profile url", "mergermarket profile", "source url", "url"},
}

ALIAS_LOOKUP = {
    alias: canonical
    for canonical, aliases in HEADER_ALIASES.items()
    for alias in aliases
}


@dataclass
class MergermarketImport:
    headers: List[str]
    raw_rows: List[List[Any]]
    deals: List[Dict[str, Any]]
    source_sheet: str
    warnings: List[str] = field(default_factory=list)


def _read_csv(data: bytes) -> Tuple[List[List[Any]], str]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)], "CSV"


def _read_xlsx(data: bytes) -> Tuple[List[List[Any]], str]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    best_rows: List[List[Any]] = []
    best_name = workbook.sheetnames[0]
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        populated = sum(1 for row in rows[:100] if any(value not in (None, "") for value in row))
        if populated > sum(1 for row in best_rows[:100] if any(value not in (None, "") for value in row)):
            best_rows = rows
            best_name = sheet.title
    return best_rows, best_name


def _find_header_row(rows: Sequence[Sequence[Any]]) -> int:
    if not rows:
        raise ValueError("The Mergermarket export is empty.")
    best_index = 0
    best_score = -1
    for index, row in enumerate(rows[:25]):
        normalised = [_normalise_header(value) for value in row]
        known = sum(1 for value in normalised if value in ALIAS_LOOKUP)
        populated = sum(1 for value in normalised if value)
        score = known * 20 + populated
        if score > best_score:
            best_index = index
            best_score = score
    return best_index


def _number(value: Any) -> Optional[float]:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    text = re.sub(r"(?i)\b(?:usd|eur|gbp|inr|mn|million|m)\b", "", text)
    text = re.sub(r"[^0-9.()\-]", "", text)
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    try:
        result = float(text)
    except ValueError:
        return None
    return -result if negative else result


def _percentage(value: Any) -> Optional[float]:
    number = _number(value)
    if number is None:
        return None
    if isinstance(value, str) and "%" in value:
        return number / 100.0
    return number / 100.0 if number > 1 else number


def _date(value: Any) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y", "%d %b %Y", "%b %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _text(value: Any) -> str:
    return str(value or "").strip()


def _dedupe_key(deal: Dict[str, Any]) -> str:
    deal_id = re.sub(r"[^a-z0-9]+", "", _text(deal.get("deal_id")).lower())
    if deal_id:
        return f"id:{deal_id}"
    parts = [
        _text(deal.get("target")).lower(),
        _text(deal.get("acquirer")).lower(),
        str(deal.get("announced_date") or ""),
    ]
    return "composite:" + "|".join(re.sub(r"[^a-z0-9]+", " ", part).strip() for part in parts)


def parse_mergermarket_export(filename: str, data: bytes) -> MergermarketImport:
    if not data:
        raise ValueError("The uploaded Mergermarket export is empty.")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        rows, sheet_name = _read_csv(data)
    elif suffix in {"xlsx", "xlsm"}:
        rows, sheet_name = _read_xlsx(data)
    else:
        raise ValueError("Upload a Mergermarket .xlsx or .csv export.")

    header_index = _find_header_row(rows)
    original_headers = [_text(value) or f"Column {index + 1}" for index, value in enumerate(rows[header_index])]
    mapped_columns: Dict[str, int] = {}
    for index, header in enumerate(original_headers):
        canonical = ALIAS_LOOKUP.get(_normalise_header(header))
        if canonical and canonical not in mapped_columns:
            mapped_columns[canonical] = index

    if "target" not in mapped_columns:
        raise ValueError("The export does not contain a recognizable Target column.")

    raw_rows = [list(row) for row in rows[header_index + 1 :] if any(value not in (None, "") for value in row)]
    deals: List[Dict[str, Any]] = []
    duplicates = 0
    seen: Dict[str, Dict[str, Any]] = {}

    def cell(row: Sequence[Any], field: str):
        index = mapped_columns.get(field)
        return row[index] if index is not None and index < len(row) else None

    for row_number, row in enumerate(raw_rows, start=header_index + 2):
        target = _text(cell(row, "target"))
        if not target:
            continue
        deal = {
            "source_row": row_number,
            "deal_id": _text(cell(row, "deal_id")),
            "announced_date": _date(cell(row, "announced_date")),
            "target": target,
            "acquirer": _text(cell(row, "acquirer")),
            "seller": _text(cell(row, "seller")),
            "deal_value_original": _number(cell(row, "deal_value_original")),
            "deal_value_original_text": _text(cell(row, "deal_value_original")),
            "deal_value_usd_mn": _number(cell(row, "deal_value_usd_mn")),
            "currency": _text(cell(row, "currency")).upper(),
            "deal_percent": _percentage(cell(row, "deal_percent")),
            "deal_status": _text(cell(row, "deal_status")),
            "deal_type": _text(cell(row, "deal_type")),
            "description": _text(cell(row, "description")),
            "target_sector": _text(cell(row, "target_sector")),
            "target_geography": _text(cell(row, "target_geography")),
            "enterprise_value_usd_mn": _number(cell(row, "enterprise_value_usd_mn")),
            "revenue_usd_mn": _number(cell(row, "revenue_usd_mn")),
            "ebitda_usd_mn": _number(cell(row, "ebitda_usd_mn")),
            "ev_ebitda_reported": _number(cell(row, "ev_ebitda_reported")),
            "ev_revenue_reported": _number(cell(row, "ev_revenue_reported")),
            "profile_url": _text(cell(row, "profile_url")),
            "relevance_tier": "Unscreened",
            "relevance_rationale": "",
        }
        key = _dedupe_key(deal)
        if key in seen:
            duplicates += 1
            existing = seen[key]
            for field, value in deal.items():
                if existing.get(field) in (None, "") and value not in (None, ""):
                    existing[field] = value
            continue
        seen[key] = deal
        deals.append(deal)

    warnings = []
    if duplicates:
        warnings.append(f"Removed {duplicates} duplicate row(s).")
    for expected in ("announced_date", "acquirer", "deal_id", "description"):
        if expected not in mapped_columns:
            warnings.append(f"No recognizable {expected.replace('_', ' ')} column was found.")

    return MergermarketImport(
        headers=original_headers,
        raw_rows=raw_rows,
        deals=deals,
        source_sheet=sheet_name,
        warnings=warnings,
    )
